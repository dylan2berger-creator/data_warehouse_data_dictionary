#!/usr/bin/env python3
"""
Build the Data Lake dataset from the source-to-data-lake mapping workbook
("BDAP Source <> Data Lake Mapping.xlsx").

The source workbook has one "Summary" sheet listing every raw ingestion
feed (source system, file type, feed/frequency/direction, the AWS Athena
table it lands in, S3 paths, description) plus one sheet per table with
its field-level mapping. See mapping_common.py for how headers are matched.

Unlike the warehouse workbook, most raw feeds in the summary sheet have no
dedicated field-mapping sheet of their own (many small XML/JSON sub-feeds
are documented once, at the catalog level, without a line-by-line field
list). Those become catalog-only entries here: real table, zero documented
fields, still worth listing so the atlas doesn't quietly drop them.

Usage:
    python3 scripts/build_datalake_dictionary.py <path-to-xlsx> [--out-json data/data_lake_dictionary.json]

Requires: openpyxl (pip install openpyxl)
Then run scripts/build_app.py to assemble index.html.
"""
import argparse
import collections
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from mapping_common import parse_detail_sheets, build_fields_and_rules, clean, norm_key

SUMMARY_SHEET_NAME = 'Summary'

# Fallback category inference for sheets/tables the catalog has no Source
# System for. Checked in order against the lowercased sheet/table name.
HEURISTIC_CATEGORY_RULES = [
    (re.compile(r'workday|^wd_|gmg'), 'Workday'),
    (re.compile(r'hertz'), 'Hertz'),
    (re.compile(r'enterprise'), 'Enterprise'),
    (re.compile(r'astech|drewtech|adas'), 'asTech / OPUS'),
    (re.compile(r'midaxo'), 'Midaxo'),
    (re.compile(r'qcomm'), 'QCOMM'),
    (re.compile(r'cushman'), 'Cushman'),
    (re.compile(r'freshservice'), 'Freshservice'),
    (re.compile(r'cultureamp'), 'CultureAmp'),
    (re.compile(r'^adp'), 'ADP'),
    (re.compile(r'navex'), 'Navex'),
    (re.compile(r'gerber|sf_rpm|sf_pricing|drp_by_location'), 'Gerber / State Farm'),
    (re.compile(r'csi'), 'CSI Survey'),
    (re.compile(r'^ca_ase|^ase_'), 'CCC - Canada (ASE)'),
    (re.compile(r'^(opportunity|assignment|repair.order|labor.assignment|purchase.order|invoice|draft.invoice|credit.memo|receipt)$'), 'CCC Cloud Connect'),
    (re.compile(r'location.service|employee.service|insurance.csi|shop.csi|checklist|jobs.detailed|profile.rates|insurance.review|insurance.carriers'), 'CCC Partner Gateway'),
    (re.compile(r'^ops.'), 'OPS'),
    (re.compile(r'dynamo'), 'Internal / Audit'),
]


def category_for(name):
    n = name.lower()
    for rx, label in HEURISTIC_CATEGORY_RULES:
        if rx.search(n):
            return label
    return 'Other'


def parse_catalog(wb):
    if SUMMARY_SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"Expected a '{SUMMARY_SHEET_NAME}' sheet, not found in workbook")
    summary_ws = wb[SUMMARY_SHEET_NAME]
    catalog = {}
    headers = [c.value for c in summary_ws[1]]
    for row in summary_ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(headers, row))
        tname = clean(d.get('AWS Athena Table Name'))
        if not tname:
            continue
        entry = {
            'source_system': clean(d.get('Source System')),
            'file_type': clean(d.get('File Type')),
            'feed_type': clean(d.get('Feed Type')),
            'file_format': clean(d.get('File Format')),
            'report_id': clean(d.get('CCC Report ID \n(if applicable)')),
            'frequency': clean(d.get('Frequency')),
            'direction': clean(d.get('Direction')),
            'processed_file_format': clean(d.get('Processed\nFile Format')),
            'processed_file_name': clean(d.get('Processed File Name')),
            'source_file_path': clean(d.get('Source File Path')),
            'target_file_path': clean(d.get('Target File Path')),
            'description': clean(d.get('Description')),
            'comments': clean(d.get('Comments')),
        }
        # A handful of Athena table names repeat with slightly different
        # catalog rows (e.g. re-listed under a second source feed); keep
        # the first, more complete, entry.
        if tname not in catalog:
            catalog[tname] = entry
    return catalog


def build_tables(catalog, sheets):
    catalog_lookup = {norm_key(k): (k, v) for k, v in catalog.items()}
    tables_out = {}
    matched_catalog_keys = set()

    for sheet_name, sheet in sheets.items():
        rows = sheet['columns']
        target_names = [r.get('target_table') for r in rows if r.get('target_table')]
        canonical = collections.Counter(target_names).most_common(1)[0][0] if target_names else sheet_name

        field_list, table_rules = build_fields_and_rules(rows)

        match = catalog_lookup.get(norm_key(canonical))
        cat_key, cat_info = match if match else (None, None)
        if cat_key:
            matched_catalog_keys.add(norm_key(cat_key))
        display_name = cat_key if cat_key else canonical
        source_system = (cat_info or {}).get('source_system')

        tables_out[sheet_name] = {
            'id': sheet_name,
            'display_name': display_name,
            'sheet_name': sheet_name,
            'category': source_system or category_for(display_name),
            'source_system': source_system,
            'file_type': (cat_info or {}).get('file_type'),
            'feed_type': (cat_info or {}).get('feed_type'),
            'file_format': (cat_info or {}).get('file_format'),
            'frequency': (cat_info or {}).get('frequency'),
            'direction': (cat_info or {}).get('direction'),
            'description': (cat_info or {}).get('description'),
            'source_file_path': (cat_info or {}).get('source_file_path'),
            'target_file_path': (cat_info or {}).get('target_file_path'),
            'catalog_comments': (cat_info or {}).get('comments'),
            'field_count': len(field_list),
            'fields': field_list,
            'table_rules': table_rules,
            'documented': True,
        }

    # Catalog rows with no field-level mapping sheet of their own still
    # describe a real data lake table -- keep them, with zero fields.
    catalog_only_count = 0
    for key, info in catalog.items():
        if norm_key(key) in matched_catalog_keys:
            continue
        catalog_only_count += 1
        tables_out[f'catalog::{key}'] = {
            'id': f'catalog::{key}',
            'display_name': key,
            'sheet_name': None,
            'category': info.get('source_system') or category_for(key),
            'source_system': info.get('source_system'),
            'file_type': info.get('file_type'),
            'feed_type': info.get('feed_type'),
            'file_format': info.get('file_format'),
            'frequency': info.get('frequency'),
            'direction': info.get('direction'),
            'description': info.get('description'),
            'source_file_path': info.get('source_file_path'),
            'target_file_path': info.get('target_file_path'),
            'catalog_comments': info.get('comments'),
            'field_count': 0,
            'fields': [],
            'table_rules': [],
            'documented': False,
        }

    return tables_out, catalog_only_count


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('xlsx_path', help='Path to the BDAP Source <> Data Lake Mapping .xlsx file')
    ap.add_argument('--out-json', default='data/data_lake_dictionary.json')
    args = ap.parse_args()

    import openpyxl
    wb = openpyxl.load_workbook(args.xlsx_path, data_only=True)

    catalog = parse_catalog(wb)
    sheets, skipped = parse_detail_sheets(wb, exclude_sheet_names={SUMMARY_SHEET_NAME})
    tables, catalog_only_count = build_tables(catalog, sheets)

    documented = sum(1 for t in tables.values() if t['documented'])
    total_fields = sum(t['field_count'] for t in tables.values())
    print(f"[datalake] Parsed {len(tables)} tables ({documented} with field-level mapping, "
          f"{catalog_only_count} catalog-only), {total_fields} fields.", file=sys.stderr)
    if skipped:
        print(f"[datalake] Skipped {len(skipped)} sheet(s) with no recognizable header: {skipped}", file=sys.stderr)

    data = {
        'id': 'datalake',
        'label': 'Data Lake',
        'tables': tables,
    }

    out_json = Path(args.out_json)
    out_json.parent.mkdir(parents=True, exist_ok=True)
    out_json.write_text(json.dumps(data, indent=1), encoding='utf-8')
    print(f"[datalake] Wrote {out_json}", file=sys.stderr)


if __name__ == '__main__':
    main()
