#!/usr/bin/env python3
"""
Build the BDAP Data Atlas from the staging-to-warehouse mapping workbook.

The source workbook has one "Summary  Index" sheet listing every warehouse
table (staging source, source system, table type, description) plus one
sheet per table containing its field-level mapping (source field -> target
field, data type, mandatory flag, transformation rule, GE data-quality
rules, notes, etc). Column headers are not fully consistent sheet to sheet,
so headers are matched by pattern rather than by fixed position.

Usage:
    python3 scripts/build_dictionary.py <path-to-xlsx> \
        [--out-json data/data_dictionary.json] \
        [--out-html index.html] \
        [--template scripts/app_template.html]

Requires: openpyxl (pip install openpyxl)
"""
import argparse
import collections
import datetime
import json
import re
import sys
from pathlib import Path

HEADER_MAP_RULES = [
    (re.compile(r'^target table \(oubound table\)$', re.I), 'target_table'),
    (re.compile(r'^target table$', re.I), 'target_table'),
    (re.compile(r'^target file$', re.I), 'target_table'),
    (re.compile(r'^target file name$', re.I), 'target_file_name'),
    (re.compile(r'^target field \(oubound table\)$', re.I), 'target_field'),
    (re.compile(r'^target field \(outbound file\)$', re.I), 'target_field_outbound'),
    (re.compile(r'^target field$', re.I), 'target_field'),
    (re.compile(r'^source table$', re.I), 'source_table'),
    (re.compile(r'^source file$', re.I), 'source_table'),
    (re.compile(r'^bi-weekly report$', re.I), 'source_table'),
    (re.compile(r'^source field.*$', re.I), 'source_field'),
    (re.compile(r'^source file column$', re.I), 'source_field'),
    (re.compile(r'^athena table$', re.I), 'athena_table'),
    (re.compile(r'^athena field$', re.I), 'athena_field'),
    (re.compile(r'^mandatory\??$', re.I), 'mandatory'),
    (re.compile(r'^data ?type( and length)?$', re.I), 'data_type'),
    (re.compile(r'^field type$', re.I), 'data_type'),
    (re.compile(r'^transformation rule$', re.I), 'transformation_rule'),
    (re.compile(r'^source format\s*/\s*examples$', re.I), 'source_format'),
    (re.compile(r'^ge ?rules?$', re.I), 'ge_rules'),
    (re.compile(r'^notes$', re.I), 'notes'),
    (re.compile(r'^comments$', re.I), 'comments'),
    (re.compile(r'^field def[ei]ni?tion$', re.I), 'field_definition'),
    (re.compile(r'^join condition$', re.I), 'join_condition'),
    (re.compile(r'^legacy col(umn)?s?$', re.I), 'legacy_column'),
    (re.compile(r'^legacy table/sp$', re.I), 'legacy_column'),
    (re.compile(r'^legacy logic/mapping$', re.I), 'legacy_logic'),
    (re.compile(r'^sp logic/mapping$', re.I), 'legacy_logic'),
    (re.compile(r'^status$', re.I), 'status'),
]

SUMMARY_SHEET_NAME = 'Summary  Index'


def norm_header(h):
    if h is None:
        return None
    h = str(h).strip()
    for rx, key in HEADER_MAP_RULES:
        if rx.match(h):
            return key
    return None


def clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v if v else None
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)):
        return v.isoformat()
    return v


def strip_suffix(name):
    return re.sub(r'\s*\([^)]*\)\s*$', '', name).strip()


def norm_key(name):
    return strip_suffix(name).strip().lower()


def category_for(name):
    n = name.lower()
    if n.startswith('dim_'):
        return 'Dimension'
    if n.startswith('fct_') or n.startswith('fact_'):
        return 'Fact'
    if n.startswith('agg_'):
        return 'Aggregate'
    if n.startswith('wd_validation') or n.startswith('wd_journal_entry'):
        return 'Workday Journal / Validation'
    if 'audit' in n:
        return 'Audit'
    return 'Other'


def parse_workbook(xlsx_path):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    if SUMMARY_SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"Expected a '{SUMMARY_SHEET_NAME}' sheet, not found in workbook")

    summary_ws = wb[SUMMARY_SHEET_NAME]
    catalog = {}
    sidx_headers = [c.value for c in summary_ws[1]]
    for row in summary_ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(sidx_headers, row))
        tname = clean(d.get('BDAP Table Name (DW)'))
        if not tname:
            continue
        catalog[tname] = {
            'staging_source': clean(d.get('Staging Source')),
            'source_system': clean(d.get('Source System')),
            'table_type': clean(d.get('Type')),
            'description': clean(d.get('BDAP Table Description (DW)')),
            'comments': clean(d.get('Comments')),
        }

    sheets = {}
    skipped = []
    for name in wb.sheetnames:
        if name == SUMMARY_SHEET_NAME:
            continue
        ws = wb[name]
        header_row_idx = None
        header_map = {}
        for r in range(1, 8):
            row_vals = [c.value for c in ws[r]]
            if not row_vals:
                continue
            mapped = {i: norm_header(v) for i, v in enumerate(row_vals) if norm_header(v)}
            if 'target_table' in mapped.values() or 'target_field' in mapped.values():
                header_row_idx = r
                header_map = mapped
                break
        if header_row_idx is None:
            skipped.append(name)
            continue

        rows_out = []
        seen_target_tables = set()
        for row in ws.iter_rows(min_row=header_row_idx + 1, values_only=True):
            rec = {}
            for i, key in header_map.items():
                if i < len(row):
                    val = clean(row[i])
                    if val is not None:
                        rec[key] = val
            if not rec:
                continue
            if rec.get('target_table'):
                seen_target_tables.add(rec['target_table'])
            if any(k in rec for k in ('target_field', 'target_field_outbound', 'source_field', 'source_table')):
                rows_out.append(rec)

        if not rows_out and not seen_target_tables:
            skipped.append(name)
            continue

        sheets[name] = {'sheet_name': name, 'columns': rows_out}

    return catalog, sheets, skipped


def build_tables(catalog, sheets):
    catalog_lookup = {norm_key(k): (k, v) for k, v in catalog.items()}
    tables_out = {}

    for sheet_name, sheet in sheets.items():
        rows = sheet['columns']
        target_names = [r.get('target_table') for r in rows if r.get('target_table')]
        canonical = collections.Counter(target_names).most_common(1)[0][0] if target_names else sheet_name

        fields = collections.OrderedDict()
        table_rules = []

        for r in rows:
            fname = r.get('target_field') or r.get('target_field_outbound')
            if not fname:
                content = r.get('source_field') or r.get('transformation_rule') or r.get('notes') or r.get('comments')
                if content and (r.get('source_table') or r.get('transformation_rule')):
                    table_rules.append({'label': r.get('source_table') or 'Note', 'text': content})
                continue
            f = fields.setdefault(fname, {
                'field_name': fname, 'data_type': None, 'mandatory': None, 'definition': None,
                'transformation_rules': [], 'ge_rules': [], 'notes': [], 'comments': [],
                'sources': [], 'join_conditions': [], 'legacy': [], 'status': None,
            })
            if r.get('data_type') and not f['data_type']:
                f['data_type'] = r['data_type']
            if r.get('mandatory'):
                if not f['mandatory'] or r['mandatory'].strip().upper().startswith('Y'):
                    f['mandatory'] = r['mandatory']
            if r.get('field_definition') and not f['definition']:
                f['definition'] = r['field_definition']
            if r.get('status') and not f['status']:
                f['status'] = r['status']
            for key, bucket in (('transformation_rule', 'transformation_rules'), ('ge_rules', 'ge_rules'),
                                ('notes', 'notes'), ('comments', 'comments'),
                                ('join_condition', 'join_conditions'), ('legacy_column', 'legacy'),
                                ('legacy_logic', 'legacy')):
                v = r.get(key)
                if v and v not in f[bucket]:
                    f[bucket].append(v)
            src = {k: r[k] for k in ('source_table', 'source_field', 'athena_table', 'athena_field', 'source_format') if r.get(k)}
            if src and src not in f['sources']:
                f['sources'].append(src)

        field_list = list(fields.values())

        match = catalog_lookup.get(norm_key(canonical))
        cat_key, cat_info = match if match else (None, None)
        display_name = cat_key if cat_key else canonical

        tables_out[sheet_name] = {
            'id': sheet_name,
            'display_name': display_name,
            'sheet_name': sheet_name,
            'category': category_for(display_name),
            'warehouse_type': (cat_info or {}).get('table_type'),
            'description': (cat_info or {}).get('description'),
            'source_system': (cat_info or {}).get('source_system'),
            'staging_source': (cat_info or {}).get('staging_source'),
            'catalog_comments': (cat_info or {}).get('comments'),
            'field_count': len(field_list),
            'fields': field_list,
            'table_rules': table_rules,
        }

    matched_keys = {norm_key(t['display_name']) for t in tables_out.values()}
    unmatched_catalog = [k for k in catalog if norm_key(k) not in matched_keys]
    return tables_out, unmatched_catalog


def render_html(data, template_path, out_html):
    template = Path(template_path).read_text(encoding='utf-8')
    if '__DATA_JSON__' not in template:
        raise SystemExit(f"Template {template_path} is missing the __DATA_JSON__ placeholder")
    data_json = json.dumps(data, separators=(',', ':')).replace('</', '<\\/')
    out = template.replace('__DATA_JSON__', data_json)
    Path(out_html).write_text(out, encoding='utf-8')


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('xlsx_path', help='Path to the BDAP Staging <> Warehouse Mapping .xlsx file')
    ap.add_argument('--out-json', default='data/data_dictionary.json')
    ap.add_argument('--out-html', default='index.html')
    ap.add_argument('--template', default='scripts/app_template.html')
    ap.add_argument('--skip-html', action='store_true', help='Only write the JSON, skip regenerating index.html')
    args = ap.parse_args()

    catalog, sheets, skipped = parse_workbook(args.xlsx_path)
    tables, unmatched_catalog = build_tables(catalog, sheets)

    total_fields = sum(t['field_count'] for t in tables.values())
    print(f"Parsed {len(tables)} tables, {total_fields} fields.", file=sys.stderr)
    if skipped:
        print(f"Skipped {len(skipped)} sheet(s) with no recognizable header: {skipped}", file=sys.stderr)
    if unmatched_catalog:
        print(f"{len(unmatched_catalog)} catalog entries had no matching sheet: {unmatched_catalog}", file=sys.stderr)

    data = {'tables': tables, 'unmatched_catalog': unmatched_catalog}

    out_json = Path(args.out_json)
    out_json.parent.mkdir(parents=True, exist_ok=True)
    out_json.write_text(json.dumps(data, indent=1), encoding='utf-8')
    print(f"Wrote {out_json}", file=sys.stderr)

    if not args.skip_html:
        render_html(data, args.template, args.out_html)
        print(f"Wrote {args.out_html}", file=sys.stderr)


if __name__ == '__main__':
    main()
